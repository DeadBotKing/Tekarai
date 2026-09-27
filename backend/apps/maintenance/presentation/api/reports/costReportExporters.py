"""CSV and Excel exporters for the maintenance cost report (time & cost).

Same conventions as the device maintenance report: UTF-8 CSV with a BOM so
Excel renders Persian correctly, and a real .xlsx workbook whose headers and
labels match the on-screen report.
"""

from __future__ import annotations

import csv
import io

from apps.maintenance.application.useCases.timeCostUseCases import (
    MaintenanceCostReportDto,
)
from apps.maintenance.presentation.api.reports import reportLabels as L

COST_COLUMNS = [
    "عنوان درخواست",
    "دستگاه",
    "واحد",
    "وضعیت",
    "تکنسین",
    "ساعت‌کار",
    "هزینه‌ی نیروی انسانی",
    "هزینه‌ی قطعات",
    "هزینه‌ی کل",
    "تاریخ ثبت",
]


def _costRow(item) -> list[object]:  # noqa: ANN001 — CostReportRowDto
    device = f"{item.deviceCode} — {item.deviceName}".strip(" —")
    return [
        item.title,
        device,
        L.departmentLabel(item.department),
        L.statusLabel(item.status),
        item.assignedToName,
        item.labourHours,
        item.labourCost,
        item.partsCost,
        item.totalCost,
        L.formatDateTime(item.createdAt),
    ]


def buildCostReportCsv(report: MaintenanceCostReportDto) -> bytes:
    """Render the report as an Excel-friendly UTF-8 CSV (BOM prefixed)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow(["گزارش هزینه‌ی نگهداری"])
    writer.writerow(["از تاریخ", report.fromDate or "—", "تا تاریخ", report.toDate or "—"])
    writer.writerow(["زمان تولید گزارش", L.formatDateTime(report.generatedAt)])
    writer.writerow([])

    writer.writerow(["خلاصه"])
    writer.writerow(["تعداد درخواست‌ها", report.workOrderCount])
    writer.writerow(["مجموع ساعت‌کار (ساعت)", report.totalLabourHours])
    writer.writerow(["مجموع هزینه‌ی نیروی انسانی", report.totalLabourCost])
    writer.writerow(["مجموع هزینه‌ی قطعات", report.totalPartsCost])
    writer.writerow(["هزینه‌ی کل نگهداری", report.totalCost])
    writer.writerow([])

    writer.writerow(["هزینه به تفکیک دستگاه"])
    writer.writerow(
        ["دستگاه", "تعداد درخواست", "ساعت‌کار", "هزینه‌ی نیروی انسانی", "هزینه‌ی قطعات", "هزینه‌ی کل"]
    )
    for group in report.byDevice:
        writer.writerow(
            [
                group.label,
                group.workOrderCount,
                group.labourHours,
                group.labourCost,
                group.partsCost,
                group.totalCost,
            ]
        )
    writer.writerow([])

    writer.writerow(["هزینه به تفکیک واحد"])
    writer.writerow(
        ["واحد", "تعداد درخواست", "ساعت‌کار", "هزینه‌ی نیروی انسانی", "هزینه‌ی قطعات", "هزینه‌ی کل"]
    )
    for group in report.byDepartment:
        writer.writerow(
            [
                L.departmentLabel(group.key),
                group.workOrderCount,
                group.labourHours,
                group.labourCost,
                group.partsCost,
                group.totalCost,
            ]
        )
    writer.writerow([])

    writer.writerow(["ساعت‌کار به تفکیک تکنسین"])
    writer.writerow(["تکنسین", "تعداد ثبت", "ساعت‌کار", "هزینه‌ی نیروی انسانی"])
    for techGroup in report.byTechnician:
        writer.writerow(
            [techGroup.technicianName, techGroup.entryCount, techGroup.hours, techGroup.cost]
        )
    writer.writerow([])

    writer.writerow(["فهرست درخواست‌ها"])
    writer.writerow(COST_COLUMNS)
    for item in report.items:
        writer.writerow(_costRow(item))

    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def buildCostReportXlsx(report: MaintenanceCostReportDto) -> bytes:
    """Render the report as a real .xlsx workbook with styled Persian blocks."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش هزینه نگهداری"
    sheet.sheet_view.rightToLeft = True

    titleFont = Font(bold=True, size=14, color="1F2937")
    headerFont = Font(bold=True, color="FFFFFF")
    labelFont = Font(bold=True, color="374151")
    headerFill = PatternFill("solid", fgColor="6D28D9")
    sectionFill = PatternFill("solid", fgColor="EDE9FE")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    rightAlign = Alignment(horizontal="right", vertical="center", wrap_text=True)

    row = 1
    titleCell = sheet.cell(row=row, column=1, value="گزارش هزینه‌ی نگهداری")
    titleCell.font = titleFont
    titleCell.alignment = rightAlign
    row += 2

    def infoRow(label1: str, value1, label2: str = "", value2="") -> None:  # noqa: ANN001
        nonlocal row
        cell1 = sheet.cell(row=row, column=1, value=label1)
        cell1.font = labelFont
        cell1.alignment = rightAlign
        sheet.cell(row=row, column=2, value=value1).alignment = rightAlign
        if label2:
            cell3 = sheet.cell(row=row, column=3, value=label2)
            cell3.font = labelFont
            cell3.alignment = rightAlign
            sheet.cell(row=row, column=4, value=value2).alignment = rightAlign
        row += 1

    def section(title: str) -> None:
        nonlocal row
        cell = sheet.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, size=12)
        cell.fill = sectionFill
        cell.alignment = rightAlign
        row += 1

    def table(headers: list[str], rows: list[list[object]]) -> None:
        nonlocal row
        for colIndex, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=colIndex, value=header)
            cell.font = headerFont
            cell.fill = headerFill
            cell.alignment = rightAlign
            cell.border = border
        row += 1
        for values in rows:
            for colIndex, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=colIndex, value=value)
                cell.alignment = rightAlign
                cell.border = border
            row += 1
        row += 1

    infoRow("از تاریخ", report.fromDate or "—", "تا تاریخ", report.toDate or "—")
    infoRow("زمان تولید گزارش", L.formatDateTime(report.generatedAt))
    row += 1

    section("خلاصه")
    infoRow("تعداد درخواست‌ها", report.workOrderCount, "مجموع ساعت‌کار", report.totalLabourHours)
    infoRow("هزینه‌ی نیروی انسانی", report.totalLabourCost, "هزینه‌ی قطعات", report.totalPartsCost)
    infoRow("هزینه‌ی کل نگهداری", report.totalCost)
    row += 1

    section("هزینه به تفکیک دستگاه")
    table(
        ["دستگاه", "تعداد درخواست", "ساعت‌کار", "هزینه‌ی نیروی انسانی", "هزینه‌ی قطعات", "هزینه‌ی کل"],
        [
            [
                group.label,
                group.workOrderCount,
                group.labourHours,
                group.labourCost,
                group.partsCost,
                group.totalCost,
            ]
            for group in report.byDevice
        ],
    )

    section("هزینه به تفکیک واحد")
    table(
        ["واحد", "تعداد درخواست", "ساعت‌کار", "هزینه‌ی نیروی انسانی", "هزینه‌ی قطعات", "هزینه‌ی کل"],
        [
            [
                L.departmentLabel(group.key),
                group.workOrderCount,
                group.labourHours,
                group.labourCost,
                group.partsCost,
                group.totalCost,
            ]
            for group in report.byDepartment
        ],
    )

    section("ساعت‌کار به تفکیک تکنسین")
    table(
        ["تکنسین", "تعداد ثبت", "ساعت‌کار", "هزینه‌ی نیروی انسانی"],
        [
            [techGroup.technicianName, techGroup.entryCount, techGroup.hours, techGroup.cost]
            for techGroup in report.byTechnician
        ],
    )

    section("فهرست درخواست‌ها")
    table(COST_COLUMNS, [_costRow(item) for item in report.items])

    widths = [30, 26, 12, 14, 16, 12, 20, 18, 18, 18]
    for colIndex, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(colIndex)].width = width

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def buildCostReportPdf(report: MaintenanceCostReportDto) -> bytes:
    """Shaped Persian A4-landscape PDF of the cost report (Phase 28).

    Carries the summary numbers in the subtitle line and the full work-order
    cost table underneath — the same columns as the CSV/XLSX exports.
    """
    from apps.maintenance.presentation.api.reports import pdfUtils

    subtitle = (
        f"بازه‌ی {report.fromDate or 'شروع'} تا {report.toDate or 'امروز'}"
        f" — {report.workOrderCount} درخواست"
        f" — ساعت‌کار: {report.totalLabourHours} ساعت"
        f" — هزینه‌ی کل: {report.totalCost}"
    )
    return pdfUtils.buildPersianTablePdf(
        title="گزارش زمان و هزینه‌ی نگهداری",
        subtitle=subtitle,
        columns=COST_COLUMNS,
        rows=[_costRow(item) for item in report.items],
        wide=True,
    )
