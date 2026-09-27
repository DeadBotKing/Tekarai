"""List-level exporters (Phase 28): work orders and the PM schedule.

Same conventions as the device/cost exporters — UTF-8 CSV with BOM, real
styled .xlsx, and a shaped Persian PDF via ``pdfUtils``.
"""

from __future__ import annotations

import csv
import io

from apps.maintenance.presentation.api.reports import pdfUtils
from apps.maintenance.presentation.api.reports import reportLabels as L
from apps.maintenance.presentation.api.reports.reportExporters import _workOrderRow

EXPORT_CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _xlsxBytes(title: str, columns: list[str], rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش"
    sheet.sheet_view.rightToLeft = True

    headerFont = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    headerFill = PatternFill("solid", fgColor="155cda")
    cellFont = Font(name="Tahoma", size=10)
    hairline = Border(
        *[Side(style="thin", color="e4e9f1")] * 3,  # left/right/top
        bottom=Side(style="thin", color="e4e9f1"),
    )
    alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    sheet.append([title])
    titleCell = sheet.cell(row=1, column=1)
    titleCell.font = Font(name="Tahoma", size=13, bold=True)
    sheet.append(columns)
    for colIndex, _ in enumerate(columns, start=1):
        cell = sheet.cell(row=2, column=colIndex)
        cell.font = headerFont
        cell.fill = headerFill
        cell.alignment = alignment
        cell.border = hairline
    for row in rows:
        sheet.append(row)
    for rowIndex in range(3, 3 + len(rows)):
        for colIndex in range(1, len(columns) + 1):
            cell = sheet.cell(row=rowIndex, column=colIndex)
            cell.font = cellFont
            cell.alignment = alignment
            cell.border = hairline
    for colIndex in range(1, len(columns) + 1):
        letter = get_column_letter(colIndex)
        maxLen = 12
        for rowIndex in range(2, min(3 + len(rows), 53)):
            value = sheet.cell(row=rowIndex, column=colIndex).value
            if value:
                maxLen = max(maxLen, min(len(str(value)) + 4, 42))
        sheet.column_dimensions[letter].width = maxLen
    sheet.freeze_panes = "A3"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _csvBytes(columns: list[str], rows: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns)
    writer.writerows(rows)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def _build(kind: str, *, title: str, columns: list[str], rows: list[list[object]], wide: bool) -> bytes:
    if kind == "xlsx":
        return _xlsxBytes(title, columns, rows)
    if kind == "pdf":
        return pdfUtils.buildPersianTablePdf(
            title=title, columns=columns, rows=rows, wide=wide
        )
    return _csvBytes(columns, rows)


# ---------------------------------------------------------------------------
# Work orders (uses the shared row/columns of the device report)
# ---------------------------------------------------------------------------


def buildWorkOrdersExport(
    orders: list,  # noqa: ANN001 — list[WorkOrderDto]
    kind: str,
) -> bytes:
    rows = [_workOrderRow(order) for order in orders]
    return _build(
        kind,
        title="فهرست درخواست‌های کار",
        columns=L.WORK_ORDER_COLUMNS,
        rows=rows,
        wide=True,
    )


# ---------------------------------------------------------------------------
# PM schedule (Phase 27 calendar data)
# ---------------------------------------------------------------------------

PM_SCHEDULE_COLUMNS: list[str] = [
    "دستگاه",
    "عنوان PM",
    "واحد",
    "مسئول",
    "زمان تخمینی (دقیقه)",
    "دوره (روز)",
    "موعد",
    "وضعیت",
]


def _pmScheduleRow(item) -> list[object]:  # noqa: ANN001 — PmScheduleItemDto
    return [
        f"{item.deviceCode} — {item.deviceName}".strip(" —"),
        item.title,
        L.departmentLabel(item.discipline),
        item.responsibleName,
        item.estimatedMinutes or "—",
        item.periodDays or "—",
        item.dueOn or "بدون تاریخ",
        "عقب‌افتاده" if item.overdue else "در برنامه",
    ]


def buildPmScheduleExport(items: list, kind: str) -> bytes:  # noqa: ANN001
    rows = [_pmScheduleRow(item) for item in items]
    return _build(
        kind,
        title="برنامه‌ی زمان‌بندی PM",
        columns=PM_SCHEDULE_COLUMNS,
        rows=rows,
        wide=False,
    )
