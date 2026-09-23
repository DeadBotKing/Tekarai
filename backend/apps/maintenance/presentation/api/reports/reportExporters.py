"""CSV and Excel exporters for the device maintenance report (Phase 23).

Both take the same ``DeviceMaintenanceReportDto`` the JSON endpoint returns and
render it with Persian headers/labels so the downloaded file matches the UI.
CSV is UTF-8 with a BOM (so Excel opens Persian correctly); XLSX is a real
workbook with a device-info block, a summary block and the work-order table.
"""

from __future__ import annotations

import csv
import io

from apps.maintenance.application.dto.maintenanceDtos import (
    DeviceMaintenanceReportDto,
)
from apps.maintenance.presentation.api.reports import reportLabels as L


def _workOrderRow(order) -> list[str]:  # noqa: ANN001 — WorkOrderDto
    return [
        order.title,
        L.typeLabel(order.orderType),
        L.priorityLabel(order.priority),
        L.statusLabel(order.status),
        L.departmentLabel(order.department),
        order.requestedByName,
        order.assignedToName,
        L.formatDateTime(order.createdAt),
        L.formatDateTime(order.slaDueAt),
        "بله" if order.overdue else "خیر",
        L.formatDateTime(order.closedAt),
        order.resolutionNote,
    ]


def buildDeviceReportCsv(report: DeviceMaintenanceReportDto) -> bytes:
    """Render the report as an Excel-friendly UTF-8 CSV (BOM prefixed)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    device = report.device
    writer.writerow(["گزارش تاریخچه‌ی نگهداری دستگاه"])
    writer.writerow(["کد دستگاه", device.code, "نام دستگاه", device.name])
    writer.writerow(
        ["محل", device.location, "واحد", L.departmentLabel(device.department)]
    )
    writer.writerow(
        ["وضعیت", L.deviceStatusLabel(device.status), "دوره‌ی PM (روز)", device.pmIntervalDays]
    )
    writer.writerow(
        ["آخرین PM", L.formatDate(device.lastPmDate), "PM بعدی", L.formatDate(device.nextDueDate)]
    )
    writer.writerow(
        ["از تاریخ", report.fromDate or "—", "تا تاریخ", report.toDate or "—"]
    )
    writer.writerow(["زمان تولید گزارش", L.formatDateTime(report.generatedAt)])
    writer.writerow([])

    summary = report.summary
    if summary is not None:
        writer.writerow(["خلاصه‌ی آماری"])
        writer.writerow(["کل درخواست‌ها", summary.totalOrders])
        writer.writerow(["باز", summary.openOrders])
        writer.writerow(["تکمیل‌شده", summary.completedOrders])
        writer.writerow(["دارای تأخیر", summary.overdueOrders])
        writer.writerow(
            ["میانگین زمان تعمیر (ساعت)", summary.mttrHours if summary.mttrHours is not None else "—"]
        )
        writer.writerow([])

    writer.writerow(["فهرست درخواست‌های کار"])
    writer.writerow(L.WORK_ORDER_COLUMNS)
    for order in report.workOrders:
        writer.writerow(_workOrderRow(order))

    # BOM so Excel detects UTF-8 and renders Persian correctly.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def buildDeviceReportXlsx(report: DeviceMaintenanceReportDto) -> bytes:
    """Render the report as a real .xlsx workbook with styled Persian blocks."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش نگهداری"
    sheet.sheet_view.rightToLeft = True

    titleFont = Font(bold=True, size=14, color="1F2937")
    headerFont = Font(bold=True, color="FFFFFF")
    labelFont = Font(bold=True, color="374151")
    headerFill = PatternFill("solid", fgColor="6D28D9")
    sectionFill = PatternFill("solid", fgColor="EDE9FE")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    rightAlign = Alignment(horizontal="right", vertical="center", wrap_text=True)

    device = report.device
    row = 1

    cell = sheet.cell(row=row, column=1, value="گزارش تاریخچه‌ی نگهداری دستگاه")
    cell.font = titleFont
    cell.alignment = rightAlign
    row += 2

    def infoRow(label1: str, value1, label2: str = "", value2="") -> None:  # noqa: ANN001
        nonlocal row
        c1 = sheet.cell(row=row, column=1, value=label1)
        c1.font = labelFont
        c1.alignment = rightAlign
        sheet.cell(row=row, column=2, value=value1).alignment = rightAlign
        if label2:
            c3 = sheet.cell(row=row, column=3, value=label2)
            c3.font = labelFont
            c3.alignment = rightAlign
            sheet.cell(row=row, column=4, value=value2).alignment = rightAlign
        row += 1

    infoRow("کد دستگاه", device.code, "نام دستگاه", device.name)
    infoRow("محل", device.location, "واحد", L.departmentLabel(device.department))
    infoRow(
        "وضعیت",
        L.deviceStatusLabel(device.status),
        "دوره‌ی PM (روز)",
        device.pmIntervalDays,
    )
    infoRow(
        "آخرین PM",
        L.formatDate(device.lastPmDate),
        "PM بعدی",
        L.formatDate(device.nextDueDate),
    )
    infoRow("از تاریخ", report.fromDate or "—", "تا تاریخ", report.toDate or "—")
    infoRow("زمان تولید گزارش", L.formatDateTime(report.generatedAt))
    row += 1

    summary = report.summary
    if summary is not None:
        sc = sheet.cell(row=row, column=1, value="خلاصه‌ی آماری")
        sc.font = Font(bold=True, size=12)
        sc.fill = sectionFill
        sc.alignment = rightAlign
        row += 1
        infoRow("کل درخواست‌ها", summary.totalOrders, "باز", summary.openOrders)
        infoRow(
            "تکمیل‌شده",
            summary.completedOrders,
            "دارای تأخیر",
            summary.overdueOrders,
        )
        infoRow(
            "میانگین زمان تعمیر (ساعت)",
            summary.mttrHours if summary.mttrHours is not None else "—",
        )
        row += 1

    tc = sheet.cell(row=row, column=1, value="فهرست درخواست‌های کار")
    tc.font = Font(bold=True, size=12)
    tc.fill = sectionFill
    tc.alignment = rightAlign
    row += 1

    headerRow = row
    for colIndex, header in enumerate(L.WORK_ORDER_COLUMNS, start=1):
        c = sheet.cell(row=headerRow, column=colIndex, value=header)
        c.font = headerFont
        c.fill = headerFill
        c.alignment = rightAlign
        c.border = border
    row += 1

    for order in report.workOrders:
        for colIndex, value in enumerate(_workOrderRow(order), start=1):
            c = sheet.cell(row=row, column=colIndex, value=value)
            c.alignment = rightAlign
            c.border = border
        row += 1

    # Reasonable column widths for the work-order table.
    widths = [30, 12, 10, 14, 12, 16, 16, 18, 18, 8, 18, 34]
    for colIndex, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(colIndex)].width = width

    sheet.freeze_panes = sheet.cell(row=headerRow + 1, column=1)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
