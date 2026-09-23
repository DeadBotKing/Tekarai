"""Phase 23 — device maintenance report + Excel/CSV/print exports.

Pins the reporting boundary the Persian frontend consumes: a per-device
maintenance-history report (device card + every work order + a statistical
summary), plus Excel (.xlsx) and CSV downloads. The printable PDF is produced
by the frontend browser-print page from the same JSON, so it is not tested
here.
"""

from __future__ import annotations

import io

from django.core.cache import cache
from django.test import TestCase
from rest_framework.test import APIClient

from tests.support.phase6Helpers import loginViaApi, seedPlatform


class MaintenanceReportBase(TestCase):
    def setUp(self) -> None:
        cache.clear()
        seedPlatform()
        self.client = APIClient()
        tokens = loginViaApi(self.client)
        self.auth = {"HTTP_AUTHORIZATION": f"Bearer {tokens['accessToken']}"}

    def createDevice(self, code: str = "PUMP-R1") -> dict:
        response = self.client.post(
            "/api/v1/maintenance/devices",
            {
                "code": code,
                "name": "پمپ خنک‌کننده",
                "location": "سالن A",
                "department": "mechanical",
                "pmIntervalDays": 30,
            },
            format="json",
            **self.auth,
        )
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()["data"]

    def submitOrder(self, deviceId: str, title: str = "عیب پمپ", priority: str = "high") -> dict:
        response = self.client.post(
            "/api/v1/maintenance/work-orders",
            {
                "deviceId": deviceId,
                "title": title,
                "description": "بررسی",
                "orderType": "corrective",
                "priority": priority,
                "requestedByName": "کاربر",
            },
            format="json",
            **self.auth,
        )
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()["data"]

    def completeOrder(self, orderId: str) -> None:
        self.assertEqual(
            self.client.post(
                f"/api/v1/maintenance/work-orders/{orderId}/assign",
                {"assignedToName": "رضا"},
                format="json",
                **self.auth,
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                f"/api/v1/maintenance/work-orders/{orderId}/status",
                {"target": "inProgress"},
                format="json",
                **self.auth,
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                f"/api/v1/maintenance/work-orders/{orderId}/status",
                {"target": "pendingApproval", "resolutionNote": "انجام شد"},
                format="json",
                **self.auth,
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                f"/api/v1/maintenance/work-orders/{orderId}/approve",
                {"note": "تأیید"},
                format="json",
                **self.auth,
            ).status_code,
            200,
        )


class DeviceReportJsonTests(MaintenanceReportBase):
    def testAuthenticationIsMandatory(self) -> None:
        device = self.createDevice()
        anonymous = APIClient()
        self.assertEqual(
            anonymous.get(f"/api/v1/maintenance/devices/{device['id']}/report").status_code,
            401,
        )

    def testReportBundlesDeviceOrdersAndSummary(self) -> None:
        device = self.createDevice()
        first = self.submitOrder(device["id"], "کار اول")
        self.submitOrder(device["id"], "کار دوم")
        self.completeOrder(first["id"])

        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report", **self.auth
        )
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]

        self.assertEqual(data["device"]["code"], device["code"])
        self.assertEqual(len(data["workOrders"]), 2)

        summary = data["summary"]
        self.assertEqual(summary["totalOrders"], 2)
        self.assertEqual(summary["completedOrders"], 1)
        self.assertEqual(summary["openOrders"], 1)
        self.assertEqual(summary["byType"]["corrective"], 2)
        self.assertIn("high", summary["byPriority"])
        # One completed order → MTTR computed (>= 0 hours).
        self.assertIsNotNone(summary["mttrHours"])

    def testEmptyDeviceReportHasZeroedSummary(self) -> None:
        device = self.createDevice("PUMP-EMPTY")
        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report", **self.auth
        )
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["workOrders"], [])
        self.assertEqual(data["summary"]["totalOrders"], 0)
        self.assertIsNone(data["summary"]["mttrHours"])

    def testFutureDateWindowExcludesOrders(self) -> None:
        device = self.createDevice()
        self.submitOrder(device["id"])
        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report",
            {"fromDate": "2099-01-01"},
            **self.auth,
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["data"]["summary"]["totalOrders"], 0)

    def testUnknownDeviceReturns404(self) -> None:
        response = self.client.get(
            "/api/v1/maintenance/devices/00000000-0000-0000-0000-000000000000/report",
            **self.auth,
        )
        self.assertEqual(response.status_code, 404, response.content)


class DeviceReportExportTests(MaintenanceReportBase):
    def testCsvExportIsUtf8WithBomAndPersianHeaders(self) -> None:
        device = self.createDevice()
        self.submitOrder(device["id"], "کار اکسپورت")

        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report",
            {"export": "csv"},
            **self.auth,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(".csv", response["Content-Disposition"])

        body = response.content
        self.assertTrue(body.startswith(b"\xef\xbb\xbf"), "CSV must start with a UTF-8 BOM")
        text = body.decode("utf-8")
        self.assertIn("گزارش تاریخچه‌ی نگهداری دستگاه", text)
        self.assertIn("فهرست درخواست‌های کار", text)
        self.assertIn("کار اکسپورت", text)

    def testXlsxExportIsAValidWorkbook(self) -> None:
        from openpyxl import load_workbook

        device = self.createDevice()
        self.submitOrder(device["id"], "کار اکسل")

        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report",
            {"export": "xlsx"},
            **self.auth,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        # The device name and a work-order row title must both be present.
        allText = "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertIn("پمپ خنک‌کننده", allText)
        self.assertIn("کار اکسل", allText)

    def testInvalidFormatIsRejected(self) -> None:
        device = self.createDevice()
        response = self.client.get(
            f"/api/v1/maintenance/devices/{device['id']}/report",
            {"export": "pdf"},
            **self.auth,
        )
        self.assertEqual(response.status_code, 400, response.content)
